from pathlib import Path
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCX = PROJECT_ROOT / "docs" / "Устав проекта AI Regulation Monitor.docx"


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


doc = Document(DOCX)
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text and any(token in text.lower() for token in ("мониторинг", "контроль", "приложение", "источник", "показател")):
        print(f"P{i}: {text}")

for ti in [5, 7]:
    table = doc.tables[ti]
    print(f"\nTABLE {ti}: rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        print(f"R{ri}: " + " | ".join(cell_text(cell)[:250] for cell in row.cells))
