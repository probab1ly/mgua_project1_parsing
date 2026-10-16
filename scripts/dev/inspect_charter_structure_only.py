from pathlib import Path

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCX = PROJECT_ROOT / "docs" / "Устав проекта AI Regulation Monitor.docx"
OUT = PROJECT_ROOT / "docs" / "_work" / "charter_current_structure.txt"


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


doc = Document(DOCX)
lines = [f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}", ""]
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        lines.append(f"P{i:03}: {text}")
lines.append("")
lines.append("TABLES")
for ti, table in enumerate(doc.tables):
    lines.append(f"T{ti}: rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows[:20]):
        lines.append(f"  R{ri}: " + " | ".join(cell_text(cell)[:220] for cell in row.cells))
    lines.append("")
OUT.write_text("\n".join(lines), encoding="utf-8")
print(OUT)
