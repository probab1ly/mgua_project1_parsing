from collections import Counter, defaultdict
from pathlib import Path
import json
import sys

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCX = PROJECT_ROOT / "docs" / "Устав проекта AI Regulation Monitor.docx"
OUT_DOC = PROJECT_ROOT / "docs" / "_work" / "charter_current_structure.txt"
OUT_SOURCES = PROJECT_ROOT / "docs" / "_work" / "active_sources_snapshot.json"

sys.path.insert(0, str(PROJECT_ROOT))
import scraper  # noqa: E402


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


doc = Document(DOCX)
lines = [f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}", ""]
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        lines.append(f"P{i:03}: {text}")
lines.append("")
lines.append("TABLES")
for ti, table in enumerate(doc.tables):
    lines.append(f"T{ti}: rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows[:15]):
        lines.append(f"  R{ri}: " + " | ".join(cell_text(cell)[:180] for cell in row.cells))
    lines.append("")
OUT_DOC.write_text("\n".join(lines), encoding="utf-8")

snapshot = []
for source in scraper.ALL_SOURCES:
    snapshot.append({
        "name": source.get("name", ""),
        "url": source.get("url", ""),
        "country": source.get("country", ""),
        "category": source.get("category", ""),
        "description": source.get("description", ""),
        "source_type": source.get("source_type", ""),
        "language": source.get("language", ""),
        "format": source.get("format", ""),
    })
OUT_SOURCES.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")

print(OUT_DOC)
print(OUT_SOURCES)
print("sources_count", len(snapshot))
print("countries", dict(Counter(item["country"] for item in snapshot)))
